"""批次1 单测：PDF 表格抽取（pdfplumber）与整篇回退。

夹具用 reportlab 造**带边框**表格 PDF（find_tables 的 lines 策略依赖页面
线条/矩形）；文本一律 ASCII——CI 无 Windows 中文字体，CJK 在 Helvetica 下
渲染为占位符，会污染断言。

无表格 PDF 用 test_chunker_headings._build_pdf 手写夹具（与既有 PDF 用例
同源），用于验证「整篇回退 → 与 pypdf 路径逐字节相等」这一承载性保证。
"""

from __future__ import annotations

import pytest

from app.agent.rag.parsers import _parse_pdf_plain, parse_document
from tests.unit.test_chunker_headings import _build_pdf


def _build_table_pdf(path, rows, *, title="Installment fee table",
                     footer="Above table is for reference only.") -> None:
    """reportlab 造带边框（GRID）表格的单页 PDF。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
    )
    table = Table(rows)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([
        Paragraph(title, styles["Heading1"]),
        Spacer(1, 6),
        table,
        Spacer(1, 6),
        Paragraph(footer, styles["BodyText"]),
    ])


# ============================================================
# 1. 带边框表格 → Markdown pipe 表
# ============================================================
def test_pdf_table_becomes_markdown(tmp_path):
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    f = tmp_path / "fee.pdf"
    _build_table_pdf(f, [
        ["Plan", "Fee"],
        ["3 months", "1.5%"],
        ["6 months", "2.5%"],
    ])

    text = parse_document(f)

    # 表头 + 分隔行 + 数据行
    assert "| Plan | Fee |" in text
    assert "| --- | --- |" in text
    assert "| 3 months | 1.5% |" in text
    assert "| 6 months | 2.5% |" in text
    # 行列不串位：3 months 与 1.5% 同行（表结构未退化成乱序词串）
    assert "| 3 months | 1.5% |" in text and "| 6 months | 2.5% |" in text
    # 文档标题保留
    assert text.startswith("# fee")
    # P1-3：该页正文含大字标题行 → 已重建真实标题层级，不再套「## 第 N 页」
    # （页码章节会让 heading_path 退化为「文档 > 第 1 页」）
    assert "## 第 1 页" not in text
    assert "# Installment fee table" in text


def test_pdf_table_text_order_preserved(tmp_path):
    """表外文本按阅读序保位：标题在前、脚注在后。"""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    f = tmp_path / "fee.pdf"
    _build_table_pdf(f, [["Plan", "Fee"], ["3 months", "1.5%"]])

    text = parse_document(f)

    assert "Installment fee table" in text
    assert "Above table is for reference only." in text
    assert text.index("Installment fee table") < text.index("| Plan | Fee |")
    assert text.index("| 3 months | 1.5% |") < text.index(
        "Above table is for reference only."
    )


def test_pdf_table_cell_pipe_escaped(tmp_path):
    """单元格内的 ``|`` 必须转义，否则列错位。"""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    f = tmp_path / "pipe.pdf"
    _build_table_pdf(f, [["Plan", "Note"], ["3 months", "A|B"]])

    text = parse_document(f)

    assert "A\\|B" in text
    # 转义后列数仍为 2：剥离转义序列后左右各一个边界符 + 中间一个分隔符
    row = next(line for line in text.splitlines() if "3 months" in line)
    assert row.replace("\\|", "").count("|") == 3


# ============================================================
# 2. 无表格 PDF → 整篇回退，与 pypdf 路径逐字节相等
# ============================================================
def test_pdf_without_tables_byte_identical(tmp_path):
    pytest.importorskip("pypdf")
    f = tmp_path / "multi.pdf"
    f.write_bytes(_build_pdf([
        "Page one about return policy.",
        "Page two about shipping.",
    ]))

    assert parse_document(f) == _parse_pdf_plain(f)
    # 同时钉死字面量，防重构漂移（页级章节结构不变）
    assert parse_document(f) == (
        "# multi\n\n## 第 1 页\n\nPage one about return policy."
        "\n\n## 第 2 页\n\nPage two about shipping."
    )


def test_pdf_empty_page_skipped_byte_identical(tmp_path):
    pytest.importorskip("pypdf")
    f = tmp_path / "gaps.pdf"
    f.write_bytes(_build_pdf(["First page.", "", "Third page."]))

    text = parse_document(f)

    assert text == _parse_pdf_plain(f)
    assert "## 第 1 页" in text and "## 第 3 页" in text
    assert "## 第 2 页" not in text


# ============================================================
# 3. 扫描件仍明确报错
# ============================================================
def test_pdf_scanned_still_errors(tmp_path):
    pytest.importorskip("pypdf")
    f = tmp_path / "scan.pdf"
    f.write_bytes(_build_pdf(["", ""]))
    with pytest.raises(ValueError, match="扫描件|无可提取文本"):
        parse_document(f)


# ============================================================
# 4. 廉价预检：无线条/矩形页不得触发 find_tables
# ============================================================
def test_pdf_precheck_skips_find_tables_without_rules(tmp_path, monkeypatch):
    """纯文本页（无 lines/rects）直接 extract_text，跳过 find_tables。

    这是 30s 解析超时的对策：find_tables 在 200 页上限内可能拖过预算，
    而超时是子进程被杀 = 上传失败，不是进程内回退。
    """
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pypdf")
    from pdfplumber.page import Page

    f = tmp_path / "plain.pdf"
    f.write_bytes(_build_pdf(["No tables on this page at all."]))

    calls: list[int] = []
    original = Page.find_tables

    def _spy(self, *args, **kwargs):
        calls.append(1)
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Page, "find_tables", _spy)

    text = parse_document(f)

    assert "No tables on this page at all." in text
    assert calls == []  # 预检命中、未进入表格检测


# ============================================================
# 5. docx 表格：表头分隔行 / 合并单元格去重 / | 转义 / 文档原序
# ============================================================
def _build_docx(path, *, merged: bool = False) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("退换货政策", level=1)
    doc.add_paragraph("支持七天无理由退货。")
    table = doc.add_table(rows=1 if merged else 2, cols=2)
    if merged:
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "时效"
    else:
        table.rows[0].cells[0].text = "时效"
        table.rows[0].cells[1].text = "7天"
        table.rows[1].cells[0].text = "运费"
        table.rows[1].cells[1].text = "商家承担"
    doc.add_heading("特殊商品", level=2)
    doc.add_paragraph("生鲜不支持七天无理由。")
    doc.save(str(path))


def test_docx_table_has_separator_and_order(tmp_path):
    pytest.importorskip("docx")
    f = tmp_path / "政策.docx"
    _build_docx(f)

    text = parse_document(f)

    assert "| 时效 | 7天 |" in text
    assert "| --- | --- |" in text
    assert "| 运费 | 商家承担 |" in text
    # 文档顺序：标题 → 正文 → 表格 → 标题 → 正文
    assert text.index("# 退换货政策") < text.index("支持七天无理由退货") \
        < text.index("| 时效 | 7天 |") < text.index("## 特殊商品") \
        < text.index("生鲜不支持")


def test_docx_merged_cell_not_duplicated(tmp_path):
    """横向合并单元格：row.cells 会重复返回同一 _tc，文本不得复制多份。"""
    pytest.importorskip("docx")
    f = tmp_path / "merged.docx"
    _build_docx(f, merged=True)

    text = parse_document(f)

    row = next(line for line in text.splitlines() if "时效" in line)
    assert row == "| 时效 |"  # 合并后只有一列
    assert text.count("时效") == 1


def test_docx_table_cell_pipe_escaped(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "时效"
    table.rows[0].cells[1].text = "A|B"
    f = tmp_path / "pipe.docx"
    doc.save(str(f))

    text = parse_document(f)

    assert "A\\|B" in text
    row = next(line for line in text.splitlines() if "时效" in line)
    assert row.replace("\\|", "").count("|") == 3


# ============================================================
# 6. HTML 表格：th/td 行列结构还原
# ============================================================
def test_html_table_becomes_pipe_rows(tmp_path):
    f = tmp_path / "p.html"
    f.write_text(
        "<html><body><h1>配送政策</h1>"
        "<table><thead><tr><th>地区</th><th>时效</th></tr></thead>"
        "<tbody><tr><td>北京</td><td>次日达</td></tr>"
        "<tr><td>偏远</td><td>不包邮</td></tr></tbody></table>"
        "<p>以上为参考时效。</p></body></html>",
        encoding="utf-8",
    )

    text = parse_document(f)

    assert "# 配送政策" in text
    assert "| 地区 | 时效 |" in text
    assert "| --- | --- |" in text
    assert "| 北京 | 次日达 |" in text
    assert "| 偏远 | 不包邮 |" in text
    # 单元格不得再被空格拼成一行（批次3 前的行为）
    assert "地区 时效" not in text
    assert "以上为参考时效。" in text


def test_html_table_implicit_cell_close(tmp_path):
    """省略 </td> / </tr> 的紧凑写法仍能正确成行（HTML 常见）。"""
    f = tmp_path / "loose.html"
    f.write_text("<table><tr><td>A<td>B<tr><td>C<td>D</table>", encoding="utf-8")

    text = parse_document(f)

    assert "| A | B |" in text
    assert "| C | D |" in text
    assert "| --- | --- |" in text


def test_html_table_cell_pipe_escaped(tmp_path):
    f = tmp_path / "pipe.html"
    f.write_text("<table><tr><td>X</td><td>A|B</td></tr></table>", encoding="utf-8")

    text = parse_document(f)

    assert "A\\|B" in text
    row = next(line for line in text.splitlines() if line.startswith("| X"))
    assert row.replace("\\|", "").count("|") == 3


def test_html_noise_filtered_and_lists_kept(tmp_path):
    """回归：script/style/head 仍被过滤，li 仍转列表项。"""
    f = tmp_path / "noise.html"
    f.write_text(
        "<html><head><title>页面标题</title>"
        "<style>.a{color:red}</style></head><body>"
        "<h2>偏远地区</h2><script>alert(1)</script>"
        "<ul><li>北京次日达</li></ul></body></html>",
        encoding="utf-8",
    )

    text = parse_document(f)

    assert "## 偏远地区" in text
    assert "- 北京次日达" in text
    assert "alert" not in text
    assert "color:red" not in text


# ============================================================
# 7. HTML caption：表格标题不得被表格状态机吞掉（Review 修正）
# ============================================================
def test_html_caption_text_kept_before_table(tmp_path):
    """caption 是表格标题：作为独立段落输出在表格之前，绝不丢内容。"""
    f = tmp_path / "caption.html"
    f.write_text(
        "<html><body><h1>费率说明</h1>"
        "<table><caption>2026年分期费率表</caption>"
        "<tr><th>期数</th><th>费率</th></tr>"
        "<tr><td>3期</td><td>0.5%</td></tr></table>"
        "<p>以上费率自 2026 年起生效。</p></body></html>",
        encoding="utf-8",
    )

    text = parse_document(f)

    assert "2026年分期费率表" in text
    assert "| 期数 | 费率 |" in text
    # caption 在表格之前（标题语义）
    assert text.index("2026年分期费率表") < text.index("| 期数 | 费率 |")
    # 表后正文不受影响
    assert "以上费率自 2026 年起生效。" in text


def test_html_unclosed_caption_not_lost(tmp_path):
    """省略 </caption> 的松散写法：表格收口时 caption 也要落盘。"""
    f = tmp_path / "loose_caption.html"
    f.write_text(
        "<table><caption>退货运费承担方对照"
        "<tr><th>场景</th><th>承担方</th></tr>"
        "<tr><td>质量问题</td><td>商家</td></tr></table>",
        encoding="utf-8",
    )

    text = parse_document(f)

    assert "退货运费承担方对照" in text
    assert "| 场景 | 承担方 |" in text
    assert "| 质量问题 | 商家 |" in text


def test_html_double_caption_keeps_both_texts(tmp_path):
    """畸形双 <caption>：两段文本都保留（宁可多一段也不丢内容）。"""
    f = tmp_path / "double_caption.html"
    f.write_text(
        "<table><caption>标题甲"
        "<caption>标题乙"
        "<tr><th>a</th></tr></table>",
        encoding="utf-8",
    )

    text = parse_document(f)

    assert "标题甲" in text
    assert "标题乙" in text
    assert "| a |" in text
    assert "页面标题" not in text
