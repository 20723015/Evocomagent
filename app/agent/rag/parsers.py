"""文档格式处理流水线（多格式接入：解析 → 结构保留 → 切分 → parent-child）。

生产知识源是 PDF/Word/Excel/HTML/飞书文档/截图。本模块提供：
- parse_document(path)：按后缀走解析器（.md/.txt 直通；.pdf 优先 Docling
  （可选依赖，未安装即回退）→ pdfplumber 抽表 + pypdf 纯文本，并叠加版式增强；
  .docx 走 python-docx（含自动编号还原）；.xlsx/.xlsm 走 openpyxl；
  .html 走内置 html.parser）→ 统一 Markdown 化文本；各格式的标题/段落/表格
  结构尽量保留成 Markdown（供标题递归切分用）；
- chunk_kb_dir(kb_dir)：目录级入口（替代原 chunk_markdown_dir），对所有
  支持后缀统一解析、切分、frontmatter 透传、parent-child 装配；
- attach_parents：parent-child 兼容装配件（新流水线已在 chunker 内联装配，
  父块 = 所属生成单元原文；本函数仅供手工构造的 Chunk 列表使用）。

parsers 侧的保真措施（2026-09 方案 P1）：
- PDF 版式增强（``_pdf_layout_hints``）：按字号重建真实标题层级、剔除跨页
  重复的页眉页脚——只做「提升为 Markdown 标题 / 整行剔除」，不改写正文抽取
  来源，因此无提示的 PDF 输出与历史实现逐字节一致；
- docx 自动编号（``_DocxNumbering``）：还原段落 numPr（含**样式继承**，本库
  语料的 "List Number" 即此形态）的十进制条款号；
- xlsx：每个工作表转 pipe 表（复用 PDF/docx/HTML 同一张表的实现）。

OCR（扫描件/图片型政策文档）为独立可选件：PaddleOCR 等重依赖不进
requirements，接入点在 parsers.parse_document 的 .pdf/.jpg 分支，
预留 parse_image_ocr(path) 钩子（Docling 的 OCR 能力经可选依赖覆盖）。
旧版 .doc 不再支持（提示转 .docx），本轮不引入 LibreOffice。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from app.agent.rag.chunker import Chunk, chunk_markdown_dir, _chunk_text
from app.agent.rag.loader import normalize_document, read_text
from app.observability.logging import get_logger

log = get_logger("app.agent.rag.parsers")

SUPPORTED_SUFFIXES = (".md", ".txt", ".pdf", ".docx", ".html", ".htm",
                      ".xlsx", ".xlsm")

# 非文档辅助文件（批次6）：同名元数据 sidecar 与 OS 垃圾文件——strict 构建
# 下也应忽略，否则未知后缀报错会把它们当坏文件中止构建。
# sidecar 必须用**文件名后缀**匹配：``Path.suffix`` 对 ``x.pdf.meta.yaml``
# 得到 ``.yaml``，会被误判为未知格式。
_OS_JUNK_NAMES = frozenset({"thumbs.db", "desktop.ini", ".ds_store"})


def _is_auxiliary_file(name: str) -> bool:
    """非文档辅助文件：同名 sidecar（``x.pdf.meta.yaml``）或 OS 垃圾文件。"""
    from app.agent.rag.governance import SIDECAR_SUFFIX

    return str(name).endswith(SIDECAR_SUFFIX) or str(name).lower() in _OS_JUNK_NAMES


def parse_document(path: Path) -> str:
    """按后缀解析为统一文本（失败抛 ValueError，由调用方跳过该文件）。

    第三方库的解析异常（pypdf PdfStreamError、zipfile BadZipFile、docx 包异常等）
    统一包装为 ValueError——上层（strict 构建）只认这一种「解析失败」语义。
    """
    suffix = path.suffix.lower()
    try:
        if suffix in (".md", ".txt"):
            return read_text(path)
        if suffix == ".pdf":
            return _parse_pdf(path)
        if suffix == ".docx":
            return _parse_docx(path)
        if suffix == ".doc":
            raise ValueError("旧版 .doc 格式不支持，请先转换为 .docx 后再上传")
        if suffix in (".html", ".htm"):
            return _parse_html(path)
        if suffix in (".xlsx", ".xlsm"):
            return _parse_xlsx(path)
        raise ValueError(f"不支持的文件格式: {suffix}（支持 {SUPPORTED_SUFFIXES}）")
    except ValueError:
        raise
    except OSError:
        raise
    except Exception as e:  # noqa: BLE001 —— 第三方解析异常统一归为「解析失败」
        raise ValueError(f"{suffix} 解析失败: {e}") from e


def _parse_pdf(path: Path) -> str:
    """逐页抽取 → 文件名一级标题 + 每页「## 第 N 页」章节（空页跳过）。

    批次1：优先 pdfplumber 按页面线条/矩形定位表格 → Markdown pipe 表，表外
    文本按阅读序交错；**全篇零表格、或抽表链路任何异常 → 整篇回退 pypdf
    extract_text()**——绝不混合双引擎输出，保证同一文档词序一致。

    P1-3 版式增强：先用 pdfplumber 的版式信息提取「标题行 + 跨页重复的页眉页脚」
    （``_pdf_layout_hints``），再叠加到既有正文抽取结果上——只做「提升为 Markdown
    标题」与「整行剔除」，不改写正文来源，因此无标题/无重复行的 PDF 输出逐字节不变。
    真实标题层级不再被「## 第 N 页」取代（跨页章节不再在页边界被切断），页眉页脚
    也不再每页重复进入正文污染 BM25 的 df 统计。

    全页无文本 = 扫描件，明确报「不支持 OCR」错误（OCR 为独立可选件；Docling
    可用时其 OCR 能力经 ``_parse_docling`` 走可选依赖）。
    """
    docling = _parse_docling(path)
    if docling is not None:
        return docling
    hints = _pdf_layout_hints(path)
    try:
        pages = _pdf_pages_with_tables(path)
    except Exception as e:  # noqa: BLE001 —— 抽表链路任何失败都整篇回退
        log.info(f"PDF 表格抽取回退纯文本（{path.name}）: {type(e).__name__}")
        pages = None
    if pages is not None:
        parts = _pdf_parts(
            path.stem, _apply_layout_hints(pages, hints),
            heading_pages=_heading_pages(hints),
        )
        if len(parts) > 1:
            return "\n\n".join(parts)
    return _parse_pdf_plain(path, hints=hints)


def _parse_pdf_plain(path: Path, hints: "_PdfLayoutHints | None" = None) -> str:
    """pypdf 纯文本抽取（批次1 前的既有实现，兼作整篇回退路径）。

    hints 缺省时自行提取（测试直接调用本函数时与 parse_document 同口径）；
    无版式提示时输出与历史实现逐字节一致。
    """
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("解析 PDF 需要 pypdf（pip install pypdf）") from e
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    if hints is None:
        hints = _pdf_layout_hints(path)
    return _join_pdf_pages(
        path.stem, _apply_layout_hints(pages, hints),
        heading_pages=_heading_pages(hints),
    )


def _parse_docling(path: Path) -> "str | None":
    """Docling 首选解析（P1-3）：Markdown 化输出（真实标题层级 + 无边框表格 + OCR）。

    可选重依赖（~200MB 级模型）：未安装（ImportError）或解析失败一律返回 None，
    由调用方回退现有解析器——绝不因第三方引擎不可用而让上传失败。输出仍走同一
    条 normalize + 标题递归切分流水线，Docling 不改变切分契约。
    """
    from app.config.settings import settings

    if not settings.rag_pdf_docling:
        return None
    try:
        from docling.document_converter import DocumentConverter
    except ImportError:
        return None  # 未接入：静默回退（可选依赖，不是错误）
    try:
        result = DocumentConverter().convert(str(path))
        markdown = result.document.export_to_markdown()
    except Exception as e:  # noqa: BLE001 —— 回退现有解析器
        log.info(f"Docling 解析失败，回退现有解析器（{path.name}）: {type(e).__name__}")
        return None
    text = (markdown or "").strip()
    if not text:
        return None
    log.info(f"Docling 解析成功（{path.name}）")
    return text


def _pdf_parts(stem: str, pages: list[str], *,
               heading_pages: "frozenset[int]" = frozenset()) -> list[str]:
    """页文本列表 → 统一 Markdown 分片（与 pypdf 路径同构）。

    heading_pages 中的页（0-based）已由版式增强重建出真实标题层级，不再套
    「## 第 N 页」页码标题——否则 heading_path 退化为「文档 > 第 3 页」，跨页
    章节的上下文价值归零。无标题的页保留页码标题（它是这些页唯一的结构信号）。
    """
    parts = [f"# {stem}"]
    for num, page_text in enumerate(pages):
        text = (page_text or "").strip()
        if not text:
            continue
        if num in heading_pages:
            parts.append(text)
        else:
            parts.append(f"## 第 {num + 1} 页\n\n{text}")
    return parts


def _join_pdf_pages(stem: str, pages: list[str], *,
                    heading_pages: "frozenset[int]" = frozenset()) -> str:
    parts = _pdf_parts(stem, pages, heading_pages=heading_pages)
    if len(parts) <= 1:
        raise ValueError("PDF 无可提取文本（扫描件走 OCR 分支）")
    return "\n\n".join(parts)


# ------------------------------------------------------------------
# P1-3：PDF 版式增强（标题层级 + 跨页重复的页眉页脚）
# ------------------------------------------------------------------
# 标题判定：字号至少比正文大 1.0pt（10.0 vs 正文 9.5 这类轻微差异不算标题）
_PDF_HEADING_MIN_DELTA = 1.0
# 只在页首/页尾 N 行内判定页眉页脚（同位置跨页重复才剔除）
_PDF_MARGIN_LINES = 2
_PDF_MAX_HEADING_LEVEL = 6


@dataclass
class _PdfLayoutHints:
    """版式提示：每页「归一化行文本 → 标题级别」+ 需整行剔除的页眉页脚文本。

    只描述「哪些行是标题/页眉」，不含正文——正文抽取来源（pdfplumber / pypdf）
    与提示来源解耦，无提示时输出与历史实现逐字节一致。
    """

    headings: dict[int, list[tuple[str, int]]] = field(default_factory=dict)
    margins: frozenset[str] = frozenset()

    @property
    def empty(self) -> bool:
        return not self.headings and not self.margins


def _pdf_norm(text: str) -> str:
    """PDF 行匹配用的归一化：去掉全部空白（抽取器之间的空格差异不影响匹配）。"""
    return re.sub(r"\s+", "", str(text or ""))


def _dominant_size(chars) -> float:
    """一行字符的主导字号（按字符数取众数，抗上下标混排）。"""
    sizes: dict[float, int] = {}
    for ch in chars or []:
        try:
            size = round(float(ch.get("size", 0.0)), 1)
        except (TypeError, ValueError):
            continue
        sizes[size] = sizes.get(size, 0) + len(str(ch.get("text", "")) or " ")
    if not sizes:
        return 0.0
    return max(sizes.items(), key=lambda kv: kv[1])[0]


def _pdf_layout_hints(path: Path) -> "_PdfLayoutHints | None":
    """pdfplumber 版式提示：标题行（按字号分级）与跨页重复的页眉页脚行。

    整篇零提示（无大字行、无重复边行）或 pdfplumber 不可用时返回 None——
    调用方据此保持正文抽取的既有输出逐字节不变。**不做数字归一化**：仅当同一位置
    跨页出现完全相同的文本才判页眉页脚，避免把「扣 3 分」/「扣 6 分」这类
    仅数字不同的正文行误删（静默丢知识是不可接受的代价）。
    """
    try:
        import pdfplumber
    except ImportError:
        return None
    try:
        with pdfplumber.open(str(path)) as pdf:
            rows_by_page: list[list[tuple[str, float]]] = []
            weight: dict[float, int] = {}
            for page in pdf.pages:
                rows: list[tuple[str, float]] = []
                for line in page.extract_text_lines():
                    text = str(line.get("text") or "").strip()
                    if not text:
                        continue
                    size = _dominant_size(line.get("chars"))
                    rows.append((text, size))
                    weight[size] = weight.get(size, 0) + len(text)
                rows_by_page.append(rows)
    except Exception as e:  # noqa: BLE001 —— 版式增强失败不影响正文抽取
        log.info(f"PDF 版式提示提取失败（{path.name}）: {type(e).__name__}")
        return None
    if not rows_by_page or not weight:
        return None

    body_size = max(weight.items(), key=lambda kv: kv[1])[0]
    big = sorted(
        (s for s in weight if s >= body_size + _PDF_HEADING_MIN_DELTA), reverse=True
    )
    level_of = {s: i + 1 for i, s in enumerate(big[:_PDF_MAX_HEADING_LEVEL])}

    headings: dict[int, list[tuple[str, int]]] = {}
    for idx, rows in enumerate(rows_by_page):
        found = [(_pdf_norm(t), level_of[s]) for t, s in rows if s in level_of]
        if found:
            headings[idx] = found
    hints = _PdfLayoutHints(headings=headings,
                            margins=_repeated_margin_lines(rows_by_page))
    return None if hints.empty else hints


def _repeated_margin_lines(rows_by_page: list[list[tuple[str, float]]]) -> frozenset[str]:
    """页首/页尾重复出现的行文本（≥ min_pages 页且占比超阈值）→ 页眉页脚。"""
    from app.config.settings import settings

    total = len(rows_by_page)
    min_pages = max(2, int(settings.rag_pdf_repeated_line_min_pages or 3))
    if total < min_pages:
        return frozenset()
    ratio = float(settings.rag_pdf_repeated_line_ratio or 0.6)
    threshold = max(min_pages, int(-(-ratio * total // 1)))  # ceil
    counts: dict[str, int] = {}
    for rows in rows_by_page:
        edge = {_pdf_norm(t) for t, _ in rows[:_PDF_MARGIN_LINES]}
        edge |= {_pdf_norm(t) for t, _ in rows[-_PDF_MARGIN_LINES:]}
        for text in edge:
            if text:
                counts[text] = counts.get(text, 0) + 1
    return frozenset(t for t, n in counts.items() if n >= threshold)


def _heading_pages(hints: "_PdfLayoutHints | None") -> frozenset[int]:
    """已重建真实标题层级的页索引（这些页不再套「## 第 N 页」页码标题）。

    页内「大字行」若全部是被判为页眉页脚的跨页重复行（真实 PDF 常见的页眉标题），
    该页并没有恢复出结构——此时保留页码标题，否则整篇 PDF 的正文会全部塌到
    根标题下、丢失页级分块边界。
    """
    if hints is None or hints.empty:
        return frozenset()
    return frozenset(
        page for page, items in hints.headings.items()
        if any(text not in hints.margins for text, _ in items)
    )


def _apply_layout_hints(pages: list[str],
                        hints: "_PdfLayoutHints | None") -> list[str]:
    """把版式提示叠加到页文本：命中行提升为 Markdown 标题 / 整行剔除页眉页脚。"""
    if hints is None or hints.empty:
        return pages
    return [
        _apply_page_hints(text, hints.headings.get(idx, []), hints.margins)
        for idx, text in enumerate(pages)
    ]


def _apply_page_hints(text: str, headings: list[tuple[str, int]],
                      margins: frozenset[str]) -> str:
    """单页叠加：行序消费标题检测（容忍个别未匹配的检测），页眉页脚整行剔除。"""
    if not headings and not margins:
        return text
    out: list[str] = []
    cursor = 0
    for line in str(text or "").split("\n"):
        stripped = line.strip()
        norm = _pdf_norm(stripped)
        if not norm:
            out.append(line)
            continue
        if norm in margins:
            continue
        level = 0
        # 有界前瞻：抽取器偶发合并/漏行时跳过失效检测，不让游标卡死
        for look in range(cursor, min(cursor + 3, len(headings))):
            if headings[look][0] == norm:
                level = headings[look][1]
                cursor = look + 1
                break
        out.append(f"{'#' * level} {stripped}" if level else line)
    return "\n".join(out)


def _pdf_pages_with_tables(path: Path) -> list[str] | None:
    """pdfplumber 逐页抽表；全篇零表格 → None（调用方整篇回退）。

    - **廉价预检**：页面无线条/矩形（纯文本页、扫描页）直接 ``extract_text``、
      跳过 ``find_tables``——``find_tables`` 在 200 页上限内可能拖过 30s 解析
      超时，而超时是子进程被杀 = 上传失败，不是进程内回退；
    - 有线页：``find_tables`` 取 bbox，表外文本用 ``crop`` 取回，按阅读序交错
      输出「表前文本 / 表 / 表间文本 / 表后文本」（多表页泛化）。
    """
    import pdfplumber

    pages: list[str] = []
    tables_found = 0
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            if not page.lines and not page.rects:
                pages.append(page.extract_text() or "")
                continue
            tables = page.find_tables()
            if not tables:
                pages.append(page.extract_text() or "")
                continue
            pages.append(_page_with_tables(page, tables))
            tables_found += len(tables)
    return pages if tables_found else None


def _page_with_tables(page, tables) -> str:
    """按阅读序交错输出页面文本与表格（表格转 Markdown pipe 表）。"""
    ordered = sorted(tables, key=lambda t: (t.bbox[1], t.bbox[0]))
    segments: list[str] = []
    cursor = 0.0
    for table in ordered:
        _, top, _, bottom = table.bbox
        if top > cursor:
            text = _band_text(page, cursor, top)
            if text:
                segments.append(text)
        markdown = _markdown_table(table.extract())
        if markdown:
            segments.append(markdown)
        cursor = max(cursor, bottom)
    if cursor < page.height:
        text = _band_text(page, cursor, page.height)
        if text:
            segments.append(text)
    return "\n\n".join(segments)


def _band_text(page, top: float, bottom: float) -> str:
    """取页面 [top, bottom) 水平带的文本；裁剪失败按无文本处理。"""
    try:
        band = page.crop(
            (0, max(top, 0.0), page.width, min(bottom, page.height)), strict=False,
        )
        return (band.extract_text() or "").strip()
    except Exception as e:  # noqa: BLE001 —— 不影响已抽出的表格
        log.info(f"PDF 文本带提取失败: {type(e).__name__}")
        return ""


_SEPARATOR_CELL = "---"


def _escape_cell(value) -> str:
    """单元格规范化：折叠空白（含内部换行）+ ``|`` 转义（防列错位）。"""
    return " ".join(str(value or "").split()).replace("|", "\\|")


def _markdown_table(rows) -> str:
    """行列表 → Markdown pipe 表（首行作表头 + 分隔行）。

    PDF / docx / HTML 三种来源共用：空单元格留空；整行皆空的行丢弃；
    列数按最宽行补齐。
    """
    if not rows:
        return ""
    cleaned = [[_escape_cell(cell) for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    header, *body = cleaned
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join([_SEPARATOR_CELL] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _parse_docx(path: Path) -> str:
    """按 XML 原始顺序处理段落与表格；Heading 1-6 → Markdown 标题。

    修复旧实现「先输出所有段落、再追加所有表格」的顺序错乱：直接遍历
    body 的 w:p / w:tbl 子元素，保持文档原序。表格转 Markdown pipe 表
    （首行作表头 + 分隔行，见 ``_docx_table_to_markdown``）。

    P1-4 自动编号还原：``Paragraph.text`` 不含 numbering.xml 的自动编号，
    政策条款的「1. 2. 3.」会静默丢失（条款引用无法检索对齐）。这里按段落的
    numPr（直接 pPr 优先，其次**样式继承**——本库语料正是用 "List Number"
    样式承载编号）还原十进制编号；非十进制（项目符号等）不生成标记。
    """
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        raise ValueError("解析 Word 需要 python-docx（pip install python-docx）") from e

    doc = Document(str(path))
    numbering = _DocxNumbering(doc)
    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            # 编号状态必须先于空段落判定推进：空段落同样打断列表运行
            number = numbering.consume(para)
            text = para.text.strip()
            if not text:
                continue
            if number and not _TYPED_ORDINAL_RE.match(text):
                text = f"{number} {text}"
            level = _heading_level(para)
            parts.append(f"{'#' * level} {text}" if level else text)
        elif child.tag == qn("w:tbl"):
            numbering.break_run()  # 表格打断列表运行
            markdown = _docx_table_to_markdown(Table(child, doc))
            if markdown:
                parts.append(markdown)
    return "\n\n".join(parts)


# ------------------------------------------------------------------
# P1-4：docx 自动编号（numbering.xml）
# ------------------------------------------------------------------

# lvlText 中的层级占位符（%1 ~ %9，%1 = ilvl 0）
_NUM_TOKEN_RE = re.compile(r"%([1-9])")
# 只还原十进制家族：其余格式（bullet / 中文数字 / 字母）不生成标记——宁简勿错
_DECIMAL_FORMATS = frozenset({"decimal", "decimalzero"})
# 段落正文已自带的序号（作者手打的「1.」「(1)」「1、」「一、」）——再加自动编号会
# 变成「1. 1. 条款」，上传链路上的双编号必须避免
_TYPED_ORDINAL_RE = re.compile(r"^\s*(?:[（(]?\d+[）).、,:：]|[一二三四五六七八九十]+[、.])")


@dataclass(frozen=True)
class _NumLevel:
    """numbering.xml 单层级定义。"""

    num_fmt: str
    lvl_text: str
    start: int


class _DocxNumbering:
    """docx 自动编号还原器（段落 numPr → 十进制编号文本）。

    词表来自 ``numbering.xml``：``w:num``（numId → abstractNumId，含
    lvlOverride/startOverride）与 ``w:abstractNum``（各 ilvl 的 numFmt /
    lvlText / start）。**样式继承**同样支持：本库语料的段落没有直接 numPr，
    编号挂在 "List Number" 样式上（w:pStyle → w:numPr）。

    计数规则：同一 numId 的连续段落为一个列表；被任何非列表段落、其它 numId
    或表格打断后从 start 重新计数（与 Word 中每个小节各自从 1 开始的观感一致），
    缩进回到上层时下层计数清零。非十进制层级返回空串（不加标记）。
    """

    def __init__(self, doc):
        self._levels: dict[str, dict[int, _NumLevel]] = {}
        self._start_overrides: dict[str, dict[int, int]] = {}
        self._abstract_of: dict[str, str] = {}
        self._counters: dict[tuple[str, int], int] = {}
        self._last_num_id: str | None = None
        self._load(doc)

    # ---------- 词表加载 ----------
    def _load(self, doc) -> None:
        from docx.oxml.ns import qn

        try:
            root = doc.part.numbering_part.element
        except Exception:  # noqa: BLE001 —— 无 numbering.xml 的文档（无自动编号）
            return
        abstract: dict[str, dict[int, _NumLevel]] = {}
        for an in root.findall(qn("w:abstractNum")):
            aid = an.get(qn("w:abstractNumId"))
            if aid is None:
                continue
            levels: dict[int, _NumLevel] = {}
            for lvl in an.findall(qn("w:lvl")):
                ilvl = _int_attr(lvl, "w:ilvl", 0)
                fmt = _child_val(lvl, "w:numFmt", "")
                text = _child_val(lvl, "w:lvlText", "")
                start = _int_attr_or(_child_val(lvl, "w:start", ""), 1)
                levels[ilvl] = _NumLevel(fmt.strip().lower(), text, start)
            abstract[aid] = levels
        for num in root.findall(qn("w:num")):
            num_id = num.get(qn("w:numId"))
            if num_id is None:
                continue
            aid = _child_val(num, "w:abstractNumId", "")
            if aid not in abstract:
                continue
            self._abstract_of[num_id] = aid
            self._levels[num_id] = abstract[aid]
            overrides = {
                _int_attr(o, "w:ilvl", 0): _int_attr_or(
                    _child_val(o, "w:startOverride", ""), 0
                )
                for o in num.findall(qn("w:lvlOverride"))
            }
            self._start_overrides[num_id] = {
                k: v for k, v in overrides.items() if v > 0
            }

    @property
    def available(self) -> bool:
        return bool(self._levels)

    # ---------- 段落消费 ----------
    def consume(self, para) -> str:
        """消费一个段落：推进编号状态并返回编号文本（无编号 → ""）。"""
        found = _paragraph_num_pr(para)
        if found is None:
            self._last_num_id = None  # 非列表段落：打断列表运行
            return ""
        num_id, ilvl = found
        if num_id not in self._levels:
            self._last_num_id = None
            return ""
        label = self._advance(num_id, ilvl)
        self._last_num_id = num_id
        return label

    def break_run(self) -> None:
        """外部元素（表格）打断列表运行。"""
        self._last_num_id = None

    def _advance(self, num_id: str, ilvl: int) -> str:
        if self._last_num_id != num_id:
            # 列表被中断（标题/正文/其它列表）→ 该 numId 重新从 start 计数
            for key in [k for k in self._counters if k[0] == num_id]:
                del self._counters[key]
        # 回到上层缩进 → 更深层级重新计数
        for key in [k for k in self._counters if k[0] == num_id and k[1] > ilvl]:
            del self._counters[key]

        level = self._levels[num_id].get(ilvl)
        if level is None or level.num_fmt not in _DECIMAL_FORMATS:
            return ""  # 非十进制（bullet/中文数字/字母）不加标记
        if not level.lvl_text:
            return ""
        start = self._start_overrides.get(num_id, {}).get(ilvl, level.start)
        current = self._counters.get((num_id, ilvl), start - 1) + 1
        self._counters[(num_id, ilvl)] = current

        def repl(match: re.Match) -> str:
            lv = int(match.group(1)) - 1
            if lv == ilvl:
                return str(current)
            return str(self._counter_of(num_id, lv))

        return _NUM_TOKEN_RE.sub(repl, level.lvl_text).strip()

    def _counter_of(self, num_id: str, ilvl: int) -> int:
        """上层层级当前计数；尚未出现时取该层 start（如 "1.1" 的首项）。"""
        value = self._counters.get((num_id, ilvl))
        if value is not None:
            return value
        level = self._levels.get(num_id, {}).get(ilvl)
        if level is None:
            return 1
        return self._start_overrides.get(num_id, {}).get(ilvl, level.start)


def _child_val(el, tag: str, default: str) -> str:
    from docx.oxml.ns import qn

    child = el.find(qn(tag))
    if child is None:
        return default
    return child.get(qn("w:val"), default) or default


def _int_attr(el, tag: str, default: int) -> int:
    from docx.oxml.ns import qn

    value = el.get(qn(tag))
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _int_attr_or(raw: str, default: int) -> int:
    try:
        return int(str(raw).strip())
    except (TypeError, ValueError):
        return default


def _num_pr_of(pPr) -> tuple[str, int] | None:
    """w:pPr 上的 (numId, ilvl)；无 numPr 返回 None。"""
    from docx.oxml.ns import qn

    if pPr is None:
        return None
    num_pr = pPr.find(qn("w:numPr"))
    if num_pr is None:
        return None
    num_id = _child_val(num_pr, "w:numId", "")
    if not num_id:
        return None
    ilvl = _int_attr_or(_child_val(num_pr, "w:ilvl", ""), 0)
    return num_id, ilvl


def _paragraph_num_pr(para) -> tuple[str, int] | None:
    """段落的编号定义：直接 pPr/numPr 优先，其次样式链（w:pStyle）继承。

    语料实测：Word 生成的清单多数把编号挂在样式上（如 "List Number"），
    段落本身没有 numPr——只查直接 numPr 会还原出 0 个编号。
    """
    from docx.oxml.ns import qn

    el = getattr(para, "_p", None)
    pPr = el.find(qn("w:pPr")) if el is not None else None
    direct = _num_pr_of(pPr)
    if direct is not None:
        return direct

    style = getattr(para, "style", None)
    seen: set[str] = set()
    while style is not None:
        style_id = getattr(style, "style_id", None) or str(id(style))
        if style_id in seen:  # basedOn 成环（损坏文档）→ 停止
            return None
        seen.add(style_id)
        style_el = getattr(style, "element", None)
        inherited = _num_pr_of(
            style_el.find(qn("w:pPr")) if style_el is not None else None
        )
        if inherited is not None:
            return inherited
        style = getattr(style, "base_style", None)
    return None


def _docx_table_to_markdown(table) -> str:
    """Word 表格 → Markdown pipe 表（首行作表头 + 分隔行）。

    批次2：横向合并的单元格在 ``row.cells`` 里会以同一 ``_tc`` 元素重复出现
    （同一文本被复制多份），按 tc 身份去重；单元格 ``|`` 由 ``_markdown_table``
    统一转义，整行皆空的行丢弃。
    """
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            cells.append(cell.text)
        if any(c.strip() for c in cells):
            rows.append(cells)
    return _markdown_table(rows)


def _heading_level(para) -> int:
    """Heading 1-6 段落样式 → Markdown 标题级别；非标题返回 0。

    python-docx 读 styles.xml 的样式名，内置标题在各语言文档中通常都是
    "Heading N"；个别本地化命名不匹配时按正文处理（宁缺勿错）。
    """
    try:
        name = (para.style.name or "").strip()
    except Exception:  # noqa: BLE001 —— 样式缺失/损坏按正文处理
        return 0
    m = re.match(r"^heading\s*(\d)$", name, re.IGNORECASE)
    if not m:
        return 0
    return min(int(m.group(1)), 6)


def _parse_html(path: Path) -> str:
    """HTML → Markdown 化正文：h1-h6 → 标题，li → 列表项，表格 → pipe 表。

    script/style/head 等非正文内容整体忽略；块内空白按 HTML 语义折叠。
    批次3：``table/thead/tbody/tr/th/td`` 独立缓冲，``tr`` 结束输出一行，
    首行后补表头分隔行；单元格 ``|`` 转义。嵌套表格不做结构还原——内层
    单元格文本并入外层命中单元格。``caption`` 是表格标题，作为独立段落
    输出在表格之前（表格模式不得吞掉表外正文）。
    """
    from html.parser import HTMLParser

    _SKIP = {"script", "style", "noscript", "template", "head", "iframe", "svg"}
    _HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
    _CELLS = {"td", "th"}
    _BLOCKS = {"p", "div", "section", "article", "aside", "header", "footer",
               "main", "nav", "table", "tr", "ul", "ol", "blockquote", "pre", "dl"}

    class _HTML2Markdown(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self._parts: list[str] = []
            self._buf: list[str] = []
            self._prefix = ""
            self._skip: dict[str, int] = {}
            # 表格状态（批次3）：仅在 depth==1 收集行列，嵌套表格文本流入外层单元格
            self._table_depth = 0
            self._rows: list[list[str]] = []
            self._row: list[str] | None = None
            self._cell: list[str] | None = None
            # caption 缓冲（Review 修正）：表格标题不进单元格，作为独立段落输出
            self._caption: list[str] | None = None

        # ---------- 纯文本 / 标题 / 列表 ----------
        def flush(self) -> None:
            if self._table_depth:
                return  # 表格模式下正文缓冲为空，交 _close_table 收口
            text = " ".join("".join(self._buf).split())
            if text:
                self._parts.append(self._prefix + text)
            self._buf = []

        def _skipping(self) -> bool:
            return any(self._skip.values())

        # ---------- 表格 ----------
        @property
        def _in_table(self) -> bool:
            return self._table_depth > 0

        def _close_cell(self) -> None:
            if self._cell is not None:
                if self._row is not None:
                    self._row.append("".join(self._cell))
                self._cell = None

        def _close_row(self) -> None:
            if self._row is not None:
                if any(cell.strip() for cell in self._row):
                    self._rows.append(self._row)
                self._row = None

        def _flush_caption(self) -> None:
            """caption 文本落为独立段落（在表格之前）；未开启时无操作。"""
            if self._caption is not None:
                text = " ".join("".join(self._caption).split())
                if text:
                    self._parts.append(text)
                self._caption = None

        def _close_table(self) -> None:
            self._close_cell()
            self._close_row()
            self._flush_caption()  # 未闭合的 caption 不丢（常见省略 </caption>）
            markdown = _markdown_table(self._rows)
            if markdown:
                self._parts.append(markdown)
            self._rows = []

        # ---------- 事件 ----------
        def handle_starttag(self, tag, attrs):
            if tag in _SKIP:
                self._skip[tag] = self._skip.get(tag, 0) + 1
                return
            if self._skipping():
                return
            if tag == "table":
                if not self._table_depth:
                    self.flush()
                self._table_depth += 1
                return
            if self._in_table:
                if self._table_depth > 1:
                    return
                if tag == "tr":
                    self._close_cell()
                    self._close_row()
                    self._row = []
                elif tag in _CELLS:
                    self._close_cell()
                    self._cell = []
                elif tag == "caption":
                    self._flush_caption()  # 畸形双 caption：先落盘前一个，不丢文本
                    self._caption = []
                return
            if tag in _HEADINGS:
                self.flush()
                self._prefix = "#" * int(tag[1]) + " "
            elif tag == "li":
                self.flush()
                self._prefix = "- "
            elif tag == "br" or tag in _BLOCKS:
                self.flush()

        def handle_endtag(self, tag):
            if tag in _SKIP:
                if self._skip.get(tag):
                    self._skip[tag] -= 1
                return
            if self._skipping():
                return
            if tag == "table":
                if self._table_depth:
                    self._table_depth -= 1
                    if not self._table_depth:
                        self._close_table()
                return
            if self._in_table:
                if self._table_depth > 1:
                    return
                if tag == "tr":
                    self._close_cell()
                    self._close_row()
                elif tag in _CELLS:
                    self._close_cell()
                elif tag == "caption":
                    self._flush_caption()
                return
            if tag in _HEADINGS or tag == "li":
                self.flush()
                self._prefix = ""
            elif tag == "br" or tag in _BLOCKS:
                self.flush()

        def handle_startendtag(self, tag, attrs):
            if tag == "br" and not self._skipping():
                if self._in_table:
                    if self._cell is not None:
                        self._cell.append(" ")
                    elif self._caption is not None:
                        self._caption.append(" ")
                else:
                    self.flush()

        def handle_data(self, data):
            if self._skipping():
                return
            if self._in_table:
                if self._cell is not None:
                    self._cell.append(data)
                elif self._caption is not None:
                    self._caption.append(data)
                return
            if not data.strip():
                return
            self._buf.append(data)

        def finish(self) -> None:
            """收口：未闭合表格先落盘，再 flush 正文缓冲。"""
            if self._table_depth:
                self._table_depth = 0
                self._close_table()
            self.flush()

    parser = _HTML2Markdown()
    parser.feed(read_text(path))
    parser.finish()
    return "\n\n".join(parser._parts)


def _parse_xlsx(path: Path) -> str:
    """Excel 工作簿 → Markdown：文件名一级标题 + 每个工作表一个二级标题 + pipe 表。

    费率表/清单/类目表类政策常以 xlsx 交付（现状 strict 构建直接报「不支持的
    格式」）。每个工作表（含表名前缀，空表跳过）转 Markdown pipe 表——复用
    PDF/docx/HTML 同一张表的 ``_markdown_table``（首行作表头 + 分隔行、单元格
    ``|`` 转义、整行皆空丢弃），因此与切分侧的表头继承（P0-4）天然对齐。

    数值格式化：``str(cell.value)``（不引入格式串），公式单元格按 openpyxl 的
    缓存值输出；合并单元格只保留左上角值（与 docx 侧去重语义一致）。
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as e:
        raise ValueError("解析 Excel 需要 openpyxl（pip install openpyxl）") from e

    try:
        # data_only=True：读公式的缓存结果（无缓存时为 None，不误报公式文本）
        workbook = load_workbook(str(path), data_only=True, read_only=True)
    except InvalidFileException as e:
        raise ValueError(f"Excel 文件无法读取: {e}") from e
    try:
        return _workbook_to_markdown(path.stem, workbook)
    finally:
        try:
            workbook.close()
        except Exception:  # noqa: BLE001 —— 关闭失败不影响已解析结果
            pass


def _workbook_to_markdown(stem: str, workbook) -> str:
    """工作簿 → Markdown（纯函数：可脱离 openpyxl 用假工作簿单测）。"""
    parts = [f"# {stem}"]
    for sheet in workbook.worksheets:
        rows = _sheet_rows(sheet)
        if not rows:
            continue
        title = str(getattr(sheet, "title", "") or "").strip()
        table = _markdown_table(rows)
        if not table:
            continue
        parts.append(f"## {title}\n\n{table}" if title else table)
    if len(parts) <= 1:
        raise ValueError("Excel 无可提取内容（所有工作表为空）")
    return "\n\n".join(parts)


def _sheet_rows(sheet) -> list[list[str]]:
    """工作表 → 行列表：全空行丢弃，行尾空单元格裁掉（避免补列把表撑宽）。"""
    rows: list[list[str]] = []
    for raw in sheet.iter_rows(values_only=True):
        cells = ["" if value is None else str(value).strip() for value in raw or ()]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            rows.append(cells)
    return rows


def parse_image_ocr(path: Path) -> str:
    """OCR 钩子（7.1）：扫描件/图片型政策文档。

    默认不实现（PaddleOCR 为可选重依赖）：接入实现放到
    requirements-ocr.txt 并在此返回识别文本；未接入抛 ValueError。
    """
    raise ValueError(
        "OCR 未接入：扫描件请安装 PaddleOCR/云服务后实现 parse_image_ocr"
    )


def attach_parents(chunks: list[Chunk], max_parent_chars: int = 4000) -> list[Chunk]:
    """parent-child 兼容装配件（手工 Chunk 列表用）。

    正常构建路径的父块已在 chunker._chunk_text 内联装配（P0-1：父块 = 所属
    生成单元原文，单元由相邻小节贪心合并而成；单节超限时回退命中位置窗口）。
    本函数只对「同 (doc, section) 多块且全部 parent_text 为空」的旧式输入
    生效；拼接结果超过 max_parent_chars 时无法保证完整包含每个子块 →
    不装配（回退命中子块），绝不截断了事。
    """
    groups: dict[tuple[str, str], list[Chunk]] = {}
    for c in chunks:
        groups.setdefault((c.doc, c.section), []).append(c)

    for group in groups.values():
        if len(group) <= 1 or any(c.parent_text for c in group):
            continue
        parent = "\n\n".join(c.text for c in group)
        if len(parent) > max_parent_chars:
            continue
        for c in group:
            c.parent_text = parent
    return chunks


def chunk_kb_dir(kb_dir: Path, *, parent_child: bool = True, strict: bool = False) -> list[Chunk]:
    """目录级多格式接入入口：所有支持后缀 → 解析 → 切分 → parent-child。

    - 隐藏目录（路径段以 `.` 开头：.staging/.trash/.git 等）一律跳过——
      KB 上传链路的暂存/隔离区绝不进入索引（v7 评审 R3）；
    - strict=True（上传/下架等**会改变正式 alias 的构建**）：任一源文件解析
      失败 → 整体抛错（含路径），杜绝「一次重建静默丢失已有知识」；
      strict=False 仅用于扫描/报告类工具（坏文件跳过，进报告人工处理）；
    - 批次6：**未知后缀在 strict 下同样报错**（原先静默 continue，表现为
      「传了 xlsx 但检索不到」）；非 strict 记录 warning 后跳过。
      辅助文件（同名 `.meta.yaml` sidecar / OS 垃圾）先行排除——生产构建
      strict=True 且 knowledge/ 下真实存在 sidecar，误判会让构建直接失败。

    .md/.txt 直接复用既有切分（含 evolved/ 沉淀 frontmatter 透传）。
    """
    chunks: list[Chunk] = []
    from app.agent.rag import governance
    from app.config.settings import settings

    enforce_meta = strict and settings.rag_doc_metadata_required
    for path in sorted(kb_dir.rglob("*")):
        if not path.is_file():
            continue
        rel = path.relative_to(kb_dir).as_posix()
        # RAG 修复计划·2：索引范围（根 + evolved/ + uploads/）与黑名单
        # （archive/.trash/.staging/隐藏/临时文件）
        if not governance.is_indexable(rel):
            continue
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            if _is_auxiliary_file(path.name):
                continue
            msg = (
                f"{rel} 不是可索引格式（支持 {', '.join(SUPPORTED_SUFFIXES)}），"
                "已跳过：请转换为受支持格式后重新上传"
            )
            if suffix == ".doc":
                msg = f"{rel} 为旧版 .doc 格式，已跳过：请先转换为 .docx 后重新上传"
            if strict:
                raise ValueError(f"strict 构建中止：{msg}")
            log.warning(msg)
            continue
        try:
            text = parse_document(path)
            normalized = normalize_document(text)
            if suffix in (".md", ".txt"):
                raw = read_text(path)
            else:
                raw = ""  # 非 md/txt：元数据走 sidecar
            # 元数据治理（RAG 修复计划·4）：所有支持格式都校验；
            # archived / external_reference 不进回答索引
            meta = governance.metadata_for(path, rel, raw)
            if enforce_meta:
                governance.validate_metadata(rel, meta)
            if meta and not governance.is_index_eligible(meta):
                continue
            if suffix in (".md", ".txt"):
                chunks.extend(
                    _chunk_text(raw, rel, kb_dir, parent_child=parent_child)
                )
            elif normalized:
                chunks.extend(
                    _chunk_text(normalized, rel, kb_dir, parent_child=parent_child)
                )
        except governance.DocumentGovernanceError:
            raise
        except (ValueError, OSError) as e:
            if strict:
                raise ValueError(
                    f"strict 构建中止：{rel} 解析失败（{e}）"
                ) from e
            # 非 strict：坏文件跳过，进报告由审核侧处理
            continue
    return chunks
